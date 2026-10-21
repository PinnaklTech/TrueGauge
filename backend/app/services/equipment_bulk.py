"""Bulk CSV/XLSX equipment import and export."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Any, Optional

from sqlalchemy.orm import Session

from app.models import EquipmentCache, new_public_id
from app.services.equipment_sync import compute_status
from app.services.qr_token import allocate_qr_token

MAX_IMPORT_ROWS = 2000

EXPORT_COLUMNS = [
    "tag",
    "name",
    "serial",
    "category",
    "manufacturer",
    "model",
    "department",
    "location",
    "status",
    "last_calibration",
    "next_calibration",
    "frequency_days",
    "owner",
    "responsible_email",
]

# Human-friendly aliases → canonical column
HEADER_ALIASES: dict[str, str] = {
    "tag": "tag",
    "tag id": "tag",
    "tag_id": "tag",
    "equipment tag": "tag",
    "name": "name",
    "equipment name": "name",
    "gauge name": "name",
    "serial": "serial",
    "serial no": "serial",
    "serial number": "serial",
    "s/n": "serial",
    "category": "category",
    "manufacturer": "manufacturer",
    "mfr": "manufacturer",
    "model": "model",
    "department": "department",
    "dept": "department",
    "location": "location",
    "status": "status",
    "last calibration": "last_calibration",
    "last_calibration": "last_calibration",
    "last cal": "last_calibration",
    "next calibration": "next_calibration",
    "next_calibration": "next_calibration",
    "next due": "next_calibration",
    "next cal": "next_calibration",
    "frequency days": "frequency_days",
    "frequency_days": "frequency_days",
    "frequency": "frequency_days",
    "owner": "owner",
    "responsible email": "responsible_email",
    "responsible_email": "responsible_email",
    "email": "responsible_email",
}

VALID_STATUSES = frozenset({"calibrated", "due-soon", "overdue", "failed", "inactive"})

TEMPLATE_EXAMPLE_ROWS = [
    {
        "tag": "EQ-001",
        "name": "Digital caliper 150mm",
        "serial": "SN-1001",
        "category": "Dimensional",
        "manufacturer": "Mitutoyo",
        "model": "CD-15AX",
        "department": "Quality",
        "location": "Lab A",
        "status": "calibrated",
        "last_calibration": "2025-01-15",
        "next_calibration": "2026-01-15",
        "frequency_days": "365",
        "owner": "Jane Doe",
        "responsible_email": "jane@example.com",
    },
    {
        "tag": "EQ-002",
        "name": "Torque wrench 50 Nm",
        "serial": "SN-2002",
        "category": "Torque",
        "manufacturer": "Snap-on",
        "model": "TECHANGLE",
        "department": "Assembly",
        "location": "Bay 3",
        "status": "due-soon",
        "last_calibration": "2025-06-01",
        "next_calibration": "2025-12-01",
        "frequency_days": "183",
        "owner": "John Smith",
        "responsible_email": "john@example.com",
    },
]


@dataclass
class RowError:
    row: int  # 1-based spreadsheet row (including header as row 1 → first data = 2)
    field: str
    message: str


@dataclass
class ParsedRow:
    row_num: int
    values: dict[str, str]
    tag: str
    name: str
    status: str
    last_calibration: Optional[date]
    next_calibration: Optional[date]
    frequency_days: int
    responsible_email: Optional[str]


@dataclass
class ImportResult:
    created: int = 0
    skipped: int = 0
    failed: int = 0
    errors: list[RowError] = field(default_factory=list)


def _norm_header(raw: str) -> str:
    key = re.sub(r"\s+", " ", (raw or "").strip().lower())
    return HEADER_ALIASES.get(key, "")


def normalize_headers(headers: list[str]) -> dict[str, int]:
    """Map canonical column → index. First occurrence wins."""
    mapping: dict[str, int] = {}
    for idx, h in enumerate(headers):
        canon = _norm_header(str(h) if h is not None else "")
        if canon and canon not in mapping:
            mapping[canon] = idx
    return mapping


def parse_optional_date(value: str) -> Optional[date]:
    raw = (value or "").strip()
    if not raw:
        return None
    # Excel sometimes emits datetime strings
    if " " in raw:
        raw = raw.split(" ", 1)[0]
    if "T" in raw:
        raw = raw.split("T", 1)[0]
    # DD/MM/YYYY or MM/DD/YYYY — prefer ISO only for v1 clarity
    try:
        return date.fromisoformat(raw[:10])
    except ValueError as exc:
        raise ValueError(f"Invalid date '{value}' (use YYYY-MM-DD)") from exc


def frequency_from_dates(last_cal: Optional[date], next_cal: Optional[date]) -> Optional[int]:
    if last_cal is None or next_cal is None:
        return None
    days = (next_cal - last_cal).days
    return days if days > 0 else None


def _cell(row: list[Any], idx: Optional[int]) -> str:
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def rows_from_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    return [[c.strip() if isinstance(c, str) else str(c or "").strip() for c in row] for row in reader]


def rows_from_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        out: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            cells: list[str] = []
            for c in row:
                if c is None:
                    cells.append("")
                elif isinstance(c, datetime):
                    cells.append(c.date().isoformat())
                elif isinstance(c, date):
                    cells.append(c.isoformat())
                elif isinstance(c, float) and c.is_integer():
                    cells.append(str(int(c)))
                else:
                    cells.append(str(c).strip())
            out.append(cells)
        return out
    finally:
        wb.close()


def parse_spreadsheet(filename: str, content: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return rows_from_xlsx(content)
    if name.endswith(".xls"):
        raise ValueError("Legacy .xls is not supported — save as .xlsx or CSV")
    # default CSV (also .csv, .txt, missing ext)
    return rows_from_csv(content)


def validate_and_parse_rows(
    sheet_rows: list[list[str]],
    *,
    existing_tags_lower: set[str],
) -> tuple[list[ParsedRow], list[RowError]]:
    errors: list[RowError] = []
    if not sheet_rows:
        errors.append(RowError(row=1, field="file", message="File is empty"))
        return [], errors

    header = sheet_rows[0]
    colmap = normalize_headers(header)
    if "name" not in colmap:
        errors.append(
            RowError(
                row=1,
                field="name",
                message="Missing required column 'name' (or 'Equipment Name')",
            )
        )
        return [], errors

    data_rows = sheet_rows[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        errors.append(
            RowError(
                row=1,
                field="file",
                message=f"Too many rows ({len(data_rows)}). Maximum is {MAX_IMPORT_ROWS}.",
            )
        )
        return [], errors

    parsed: list[ParsedRow] = []
    seen_tags: set[str] = set()

    for offset, raw in enumerate(data_rows):
        row_num = offset + 2  # header is row 1
        if not raw or all(not str(c).strip() for c in raw):
            continue

        def get(col: str) -> str:
            return _cell(raw, colmap.get(col))

        name = get("name")
        if not name:
            errors.append(RowError(row=row_num, field="name", message="Name is required"))
            continue

        serial = get("serial")
        tag_raw = get("tag")
        tag = tag_raw or serial or name[:32]
        tag = tag.strip()
        if not tag:
            errors.append(RowError(row=row_num, field="tag", message="Could not derive a tag"))
            continue

        tag_key = tag.casefold()
        if tag_key in seen_tags:
            errors.append(
                RowError(row=row_num, field="tag", message=f"Duplicate tag in file: {tag}")
            )
            continue
        if tag_key in existing_tags_lower:
            errors.append(
                RowError(
                    row=row_num,
                    field="tag",
                    message=f"Tag already exists — skipped: {tag}",
                )
            )
            continue

        status_raw = get("status").lower().replace("_", "-").replace(" ", "-")
        if not status_raw:
            status = "calibrated"
        elif status_raw in VALID_STATUSES:
            status = status_raw
        else:
            errors.append(
                RowError(
                    row=row_num,
                    field="status",
                    message=f"Invalid status '{get('status')}'",
                )
            )
            continue

        try:
            last_cal = parse_optional_date(get("last_calibration"))
            next_cal = parse_optional_date(get("next_calibration"))
        except ValueError as exc:
            errors.append(RowError(row=row_num, field="date", message=str(exc)))
            continue

        if last_cal is not None and next_cal is not None and next_cal <= last_cal:
            errors.append(
                RowError(
                    row=row_num,
                    field="next_calibration",
                    message="Next calibration must be after last calibration",
                )
            )
            continue

        freq = frequency_from_dates(last_cal, next_cal)
        if freq is None:
            freq_raw = get("frequency_days")
            if freq_raw:
                try:
                    freq = int(float(freq_raw))
                    if freq < 1:
                        raise ValueError("frequency must be >= 1")
                except ValueError:
                    errors.append(
                        RowError(
                            row=row_num,
                            field="frequency_days",
                            message=f"Invalid frequency_days '{freq_raw}'",
                        )
                    )
                    continue
            else:
                freq = 365

        if status == "inactive" and next_cal is not None:
            status = compute_status(next_cal)

        email = get("responsible_email") or None
        if email and "@" not in email:
            errors.append(
                RowError(
                    row=row_num,
                    field="responsible_email",
                    message=f"Invalid email '{email}'",
                )
            )
            continue

        seen_tags.add(tag_key)
        parsed.append(
            ParsedRow(
                row_num=row_num,
                values={
                    "tag": tag,
                    "name": name,
                    "serial": serial,
                    "category": get("category"),
                    "manufacturer": get("manufacturer"),
                    "model": get("model"),
                    "department": get("department"),
                    "location": get("location"),
                    "owner": get("owner"),
                },
                tag=tag,
                name=name,
                status=status,
                last_calibration=last_cal,
                next_calibration=next_cal,
                frequency_days=freq,
                responsible_email=email,
            )
        )

    return parsed, errors


def build_template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for row in TEMPLATE_EXAMPLE_ROWS:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")


def build_export_csv(rows: list[EquipmentCache]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for r in rows:
        writer.writerow(
            {
                "tag": r.tag or "",
                "name": r.name or "",
                "serial": r.serial or "",
                "category": r.category or "",
                "manufacturer": r.manufacturer or "",
                "model": r.model or "",
                "department": r.department or "",
                "location": r.location or "",
                "status": r.status or "",
                "last_calibration": r.last_calibration.isoformat() if r.last_calibration else "",
                "next_calibration": r.next_calibration.isoformat() if r.next_calibration else "",
                "frequency_days": str(r.frequency_days or 365),
                "owner": r.owner or "",
                "responsible_email": r.responsible_email or "",
            }
        )
    return buf.getvalue().encode("utf-8-sig")


def existing_tags_for_tenant(db: Session, tenant_id: int) -> set[str]:
    rows = (
        db.query(EquipmentCache.tag)
        .filter(EquipmentCache.tenant_id == tenant_id)
        .all()
    )
    return {(t or "").strip().casefold() for (t,) in rows if (t or "").strip()}


def import_equipment_rows(
    db: Session,
    *,
    tenant_id: int,
    parsed: list[ParsedRow],
) -> int:
    """Insert validated rows. Returns created count."""
    created = 0
    now = datetime.now(timezone.utc)
    for item in parsed:
        row = EquipmentCache(
            tenant_id=tenant_id,
            public_id=new_public_id(),
            qr_token=allocate_qr_token(
                db, model_cls=EquipmentCache, column_attr=EquipmentCache.qr_token
            ),
            odoo_id=None,
            source="local",
            tag=item.tag,
            name=item.name,
            category=item.values.get("category", ""),
            manufacturer=item.values.get("manufacturer", ""),
            model=item.values.get("model", ""),
            serial=item.values.get("serial", ""),
            department=item.values.get("department", ""),
            location=item.values.get("location", ""),
            status=item.status,
            last_calibration=item.last_calibration,
            next_calibration=item.next_calibration,
            frequency_days=item.frequency_days,
            owner=item.values.get("owner", ""),
            responsible_email=item.responsible_email,
            raw_payload=None,
            synced_at=now,
        )
        db.add(row)
        created += 1
    return created


def run_import(
    db: Session,
    *,
    tenant_id: int,
    filename: str,
    content: bytes,
) -> ImportResult:
    result = ImportResult()
    try:
        sheet_rows = parse_spreadsheet(filename, content)
    except ValueError as exc:
        result.failed = 1
        result.errors.append(RowError(row=1, field="file", message=str(exc)))
        return result
    except Exception as exc:  # noqa: BLE001 — surface parse failures to client
        result.failed = 1
        result.errors.append(
            RowError(row=1, field="file", message=f"Could not read file: {exc}")
        )
        return result

    existing = existing_tags_for_tenant(db, tenant_id)
    parsed, errors = validate_and_parse_rows(sheet_rows, existing_tags_lower=existing)

    # Classify: skip vs fail
    for err in errors:
        if "already exists" in err.message.lower() or "duplicate tag in file" in err.message.lower():
            result.skipped += 1
        else:
            result.failed += 1
    result.errors = errors

    if not parsed:
        return result

    created = import_equipment_rows(db, tenant_id=tenant_id, parsed=parsed)
    result.created = created
    return result
